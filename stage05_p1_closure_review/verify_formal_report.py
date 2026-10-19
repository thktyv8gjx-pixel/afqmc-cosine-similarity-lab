"""Read-only formula/source/artifact gate. Run with document workspace Python.

Never writes a report or imports an AFQMC evaluator. Full source/identity remain
in local_only; public outputs contain only hashes, counts and check outcomes.
"""
import hashlib,json,re,unicodedata
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from pypdf import PdfReader
R=Path(__file__).resolve().parents[1];O=Path(__file__).parent;L=O/'local_only'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def flat(t):return re.sub(r'\s','',unicodedata.normalize('NFKC',t))
def check_source(source):
    s=flat(source)
    assert r'\operatorname{cos}(x_i,y_i)' not in s and 'cos(x_i,y_i)' not in s
    assert r'y_i\in\{0,1\}' in s
    assert r'x_i^{(1)}=\phi(N(s_i^{(1)}))' in s and r'x_i^{(2)}=\phi(N(s_i^{(2)}))' in s
    assert r'\haty_i=\mathbb{1}[\operatorname{cos}(x_i^{(1)},x_i^{(2)})\ge\tau]' in s
    return True
def main():
    registry=read(O/'formal_report_registry.json');source=R/registry['canonical_source']['path']
    assert sha(source)==registry['canonical_source']['sha256'];check_source(source.read_text(encoding='utf-8'))
    assert sha(R/registry['regeneration_adapter'])==registry['regeneration_adapter_sha256']
    assert sha(R/registry['formatting_helper']['path'])==registry['formatting_helper']['sha256']
    for name in ['stage04_report_draft/report_body.md.in','stage04_report_draft/报告正文初稿_公开安全版.md']:
        historical=(R/name).read_text(encoding='utf-8');assert '禁止作为最终报告再生成源' in historical.splitlines()[0]
        try:check_source(historical)
        except AssertionError:pass
        else:raise AssertionError('Historical flawed source must fail the formal-source gate')
    original=R/registry['canonical_source']['copied_from'];assert source.read_bytes()==original.read_bytes()
    for x in registry['figure_sources']:assert sha(R/x['path'])==x['sha256']
    pairdir=R/registry['report_pair_directory'];doc=next(pairdir.glob('*_实验报告1_*.docx'));pdf=doc.with_suffix('.pdf')
    assert sha(doc)==registry['docx_sha256'] and sha(pdf)==registry['pdf_sha256']
    d=Document(doc);reader=PdfReader(pdf);assert len(reader.pages)==15
    text='\n'.join(p.extract_text() for p in reader.pages);ft=flat(text)
    identity=read(R/'stage04_final_report/revision_highest_audit/local_only/identity.json')
    for row,value in zip(d.tables[0].rows[2:7],identity['values']):assert row.cells[1].text==value and flat(value) in ft
    assert '【待确认：实际提交日期】'==d.tables[0].rows[7].cells[1].text
    assert '通信工程学院' in ft
    with ZipFile(doc) as z:
        root=E.fromstring(z.read('word/document.xml'));maths=[''.join(m.xpath('.//m:t/text()',namespaces=NS)) for m in root.findall('.//m:oMath',NS)]
        assert not any('cos(xi,yi)' in flat(x) for x in maths)
        assert any('cos(xi(1),xi(2))≥τ' in flat(x) for x in maths)
        assert any('yi∈{0,1}' in flat(x) for x in maths)
        assert any('xi(1)=ϕ(N(si(1)))' in x and 'xi(2)=ϕ(N(si(2)))' in x for x in maths)
        assert 'cos(xi,yi)' not in ft and 'cos(xi(1),xi(2))≥τ' in ft
        assert len(maths)==52
    headings=[p.text for p in d.paragraphs]
    assert all(any(p.startswith(str(i)+' ') for p in headings) for i in range(1,5))
    assert sum(bool(re.match(r'^[23]\.\d ',p)) for p in headings)==7
    assert 'AI工具使用说明' in headings
    gate=read(R/'stage04_final_report/revision_highest_audit/document_gate.json')
    for p in gate['pages_evidence']:
        assert sha(pairdir/f'word_render/page-{p["page"]}.png')==p['word_render_sha256']
        assert sha(pairdir/f'pdf_render/page-{p["page"]:02d}.png')==p['pdf_render_sha256']
    result={'passed':True,'formal_source_sha256':sha(source),'docx_sha256':sha(doc),'pdf_sha256':sha(pdf),'pages':15,'native_math_objects':52,'formula_role_checks_source_docx_pdf':'PASS','confirmed_five_identity_fields':'PASS','date_pending':True,'original_institution_retained':True,'report_pair_rewritten':False,'prior_fifteen_page_visual_evidence_bound_by_hash':True,'historical_draft_expressions_are_not_formal_sources':True}
    (O/'formal_report_check.json').write_bytes((json.dumps(result,indent=2)+'\n').encode())
    print(json.dumps(result))
if __name__=='__main__':main()
